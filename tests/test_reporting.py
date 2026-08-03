"""Tests for repogen.reporting - export, combine_results, html_report."""

from __future__ import annotations

import html.parser
import json
import struct
import zlib
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from repogen.reporting.combine_results import (
    RESULT_FILE_MAP,
    CombinedResults,
    combine_results,
)
from repogen.reporting.export import (
    VALID_RESULT_TYPES,
    _EXCEL_MAX_CELL_CHARS,
    _build_summary,
    _convert_numpy_types,
    _dataframe_to_records,
    _sanitize_excel_str,
    export_combined,
    export_results,
)
from repogen.reporting.html_report import (
    _embed_figure,
    _render_table,
    generate_html_report,
)

# ---------------------------------------------------------------------------
# Synthetic test data fixtures
# ---------------------------------------------------------------------------

TEST_FILENAMES: dict[str, str] = {
    "gene": "{study}_gene_results.parquet",
    "pathway": "{study}_pathway_results.parquet",
    "drug": "{study}_drug_enrichment.parquet",
    "atc": "{study}_atc_enrichment_results.parquet",
    "spredixcan": "{study}_spredixcan_meta_analysis.parquet",
    "spredixcan_per_tissue": "{study}_spredixcan_per_tissue.parquet",
    "correlation": "{study}_drug_summary.parquet",
    "correlation_per_tissue": "{study}_per_tissue_results.parquet",
    "mr": "{study}_mr_results.parquet",
    "mr_drugs": "{study}_mr_drug_matches.parquet",
    "mr_verdicts": "{study}_mr_target_verdicts.parquet",
}


@pytest.fixture()
def gene_results() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "gene_symbol": ["GENE1", "GENE2", "GENE3", "GENE4", "GENE5"],
            "chr": [1, 1, 2, 6, 6],
            "start": [1000, 2000, 3000, 28000000, 29000000],
            "magma_z": [4.5, 3.2, 2.1, 5.0, 1.5],
            "magma_p": [3.4e-6, 6.8e-4, 0.02, 2.8e-7, 0.07],
            "fdr_q": [1.7e-5, 0.002, 0.03, 7e-7, 0.09],
            "n_snps": [120, 85, 45, 200, 30],
            "biotype": ["protein_coding"] * 5,
            "in_mhc": [False, False, False, True, True],
            "annotation_mode": ["proximity"] * 5,
        }
    )


@pytest.fixture()
def pathway_results() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "pathway_name": ["PathA", "PathB", "PathC", "PathD", "PathE"],
            "source_db": ["GO_BP", "KEGG", "REACTOME", "GO_BP", "WP"],
            "p_value": [1e-5, 0.001, 0.02, 0.04, 0.08],
            "fdr_q": [5e-5, 0.005, 0.04, 0.06, 0.12],
            "n_genes_in_set": [50, 30, 80, 20, 15],
            "beta": [0.5, 0.3, 0.2, 0.15, 0.1],
            "driving_genes": [
                ["G1", "G2"],
                ["G3"],
                ["G4", "G5", "G6"],
                [],
                ["G7"],
            ],
        }
    )


@pytest.fixture()
def drug_enrichment() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "drug_name": ["DrugA", "DrugB", "DrugC", "DrugD", "DrugE"],
            "drug_chembl_id": ["CHEMBL1", "CHEMBL2", "CHEMBL3", "CHEMBL4", "CHEMBL5"],
            "magma_p": [1e-6, 0.001, 0.03, 0.05, 0.1],
            "magma_fdr_q": [5e-6, 0.005, 0.04, 0.08, 0.2],
            "magma_beta": [0.8, 0.5, 0.3, 0.2, 0.1],
            "max_phase": [4, 3, 2, 1, 0],
            "atc_codes": [
                ["N05A", "N06A"],
                ["N05B"],
                [],
                ["C07A"],
                None,
            ],
            "target_genes": [
                json.dumps(["TG1", "TG2"]),
                json.dumps(["TG3"]),
                json.dumps([]),
                json.dumps(["TG4"]),
                json.dumps(["TG5"]),
            ],
        }
    )


@pytest.fixture()
def atc_enrichment() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "atc_code": ["N05A", "N06A", "C07A"],
            "atc_description": ["Antipsychotics", "Antidepressants", "Beta blockers"],
            "atc_level": [3, 3, 3],
            "gls_p": [1e-4, 0.01, 0.05],
            "gls_fdr": [5e-4, 0.03, 0.08],
            "n_drugs": [15, 12, 8],
            "contributing_drugs": ["D1, D2, D3", "D4, D5", "D6"],
        }
    )


@pytest.fixture()
def neg_correlation_summary() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "drug_name": ["DrugA", "DrugX", "DrugY", "DrugZ", "DrugW"],
            "drug_chembl_id": ["CHEMBL1", "CHEMBL10", "CHEMBL11", None, None],
            "best_spearman_rho": [-0.35, -0.28, -0.22, -0.15, -0.10],
            "best_tissue": [
                "Brain_Cortex",
                "Brain_Hippocampus",
                "Brain_Cortex",
                "Brain_Cortex",
                "Brain_Cortex",
            ],
            "n_tissues_fdr_significant": [5, 3, 1, 0, 0],
        }
    )


@pytest.fixture()
def mr_results() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "gene_symbol": ["GENA", "GENB", "GENC"],
            "eqtl_source": ["eQTLGen", "MetaBrain", "eQTLGen"],
            "mr_beta": [0.15, -0.22, 0.08],
            "mr_se": [0.03, 0.05, 0.04],
            "mr_pval": [1e-6, 1e-5, 0.04],
            "mr_significant": [True, True, False],
            "confidence_tier": ["high", "high", "low"],
            "n_instruments": [8, 5, 3],
            "pp_h4": [0.95, 0.88, 0.3],
        }
    )


@pytest.fixture()
def mr_drug_matches() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "gene_symbol": ["GENA", "GENA", "GENB"],
            "drug_name": ["DrugA", "DrugM", "DrugN"],
            "drug_chembl_id": ["CHEMBL1", "CHEMBL20", "CHEMBL21"],
            "confidence_tier": ["high", "medium", "high"],
            "direction_concordant": [True, False, True],
            "interaction_type": ["inhibitor", "agonist", "antagonist"],
            "max_phase": [4, 2, 3],
        }
    )


@pytest.fixture()
def spredixcan_meta() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "gene_symbol": ["SPG1", "SPG2", "SPG3", "SPG4"],
            "meta_zscore": [4.0, -3.5, 2.0, 1.2],
            "meta_pvalue": [3e-5, 2e-4, 0.04, 0.2],
        }
    )


@pytest.fixture()
def all_fixtures(
    gene_results,
    pathway_results,
    drug_enrichment,
    atc_enrichment,
    neg_correlation_summary,
    mr_results,
    mr_drug_matches,
    spredixcan_meta,
):
    return {
        "gene": gene_results,
        "pathway": pathway_results,
        "drug": drug_enrichment,
        "atc": atc_enrichment,
        "spredixcan": spredixcan_meta,
        "correlation": neg_correlation_summary,
        "mr": mr_results,
        "mr_drugs": mr_drug_matches,
    }


def _write_test_results(tmp_path, study_name, result_types, fixtures):
    """Write synthetic Parquet files to the expected directory structure."""
    for rt in result_types:
        info = RESULT_FILE_MAP[rt]
        subdir = tmp_path / study_name / info["subdir"]
        subdir.mkdir(parents=True, exist_ok=True)
        filename = TEST_FILENAMES[rt].format(study=study_name)
        fixtures[rt].to_parquet(subdir / filename)


def _create_dummy_png(path: Path) -> None:
    """Write a minimal valid PNG file (1x1 pixel, transparent)."""
    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
    ihdr_crc = struct.pack(">I", zlib.crc32(b"IHDR" + ihdr) & 0xFFFFFFFF)
    idat_data = zlib.compress(b"\x00\x00\x00\x00\x00")
    idat_crc = struct.pack(">I", zlib.crc32(b"IDAT" + idat_data) & 0xFFFFFFFF)
    iend_crc = struct.pack(">I", zlib.crc32(b"IEND") & 0xFFFFFFFF)
    png = (
        sig
        + struct.pack(">I", 13)
        + b"IHDR"
        + ihdr
        + ihdr_crc
        + struct.pack(">I", len(idat_data))
        + b"IDAT"
        + idat_data
        + idat_crc
        + struct.pack(">I", 0)
        + b"IEND"
        + iend_crc
    )
    path.write_bytes(png)


# ===================================================================
# TestExportResults
# ===================================================================


class TestExportResults:
    """Per-module export tests."""

    def test_export_csv_creates_file(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["csv"],
        )
        assert len(paths) == 1
        assert paths[0].exists()
        assert paths[0].name == "TEST_gene.csv"
        df = pd.read_csv(paths[0])
        assert len(df) == 5

    def test_export_json_creates_file(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["json"],
        )
        assert len(paths) == 1
        data = json.loads(paths[0].read_text())
        assert data["result_type"] == "gene"
        assert data["study"] == "TEST"
        assert "timestamp" in data
        assert "summary" in data
        assert "data" in data

    def test_export_json_summary_fields(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["json"],
        )
        data = json.loads(paths[0].read_text())
        assert "n_genes_tested" in data["summary"]
        assert data["summary"]["n_genes_tested"] == 5
        assert data["summary"]["n_significant"] == 4

    def test_export_csv_list_columns_stringified(self, tmp_path, pathway_results):
        parquet_path = tmp_path / "pathways.parquet"
        pathway_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="pathway",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["csv"],
        )
        df = pd.read_csv(paths[0])
        first_dg = df["driving_genes"].iloc[0]
        parsed = json.loads(first_dg)
        assert isinstance(parsed, list)
        assert parsed == ["G1", "G2"]

    def test_export_invalid_result_type_raises(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        with pytest.raises(ValueError, match="Unknown result_type"):
            export_results(
                result_type="invalid",
                results_path=parquet_path,
                output_dir=tmp_path / "out",
            )

    def test_export_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            export_results(
                result_type="gene",
                results_path=tmp_path / "nonexistent.parquet",
                output_dir=tmp_path / "out",
            )

    def test_export_with_metadata(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        meta_path = tmp_path / "meta.json"
        meta_path.write_text(json.dumps({"parameters": {"window": 35}}))

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            metadata_path=meta_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["json"],
        )
        data = json.loads(paths[0].read_text())
        assert data["parameters"]["window"] == 35

    def test_export_empty_dataframe(self, tmp_path):
        empty = pd.DataFrame({"gene_symbol": [], "fdr_q": []})
        parquet_path = tmp_path / "empty.parquet"
        empty.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["csv", "json"],
        )
        assert len(paths) == 2
        for p in paths:
            assert p.exists()


# ===================================================================
# TestExportResultsXlsx - per-result XLSX and truncation
# ===================================================================


class TestExportResultsXlsx:
    """Per-result XLSX export and cell-length truncation tests."""

    def test_xlsx_format_creates_file(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["xlsx"],
        )
        assert len(paths) == 1
        assert paths[0].exists()
        assert paths[0].name == "TEST_gene.xlsx"
        df = pd.read_excel(paths[0])
        assert len(df) == 5
        assert "gene_symbol" in df.columns

    def test_mixed_formats_csv_json_xlsx(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["csv", "json", "xlsx"],
        )
        assert len(paths) == 3
        extensions = {p.suffix for p in paths}
        assert extensions == {".csv", ".json", ".xlsx"}

    def test_default_formats_no_xlsx(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
        )
        assert len(paths) == 2
        extensions = {p.suffix for p in paths}
        assert extensions == {".csv", ".json"}

    def test_xlsx_truncation_applied(self, tmp_path):
        long_val = "x" * 50000
        df = pd.DataFrame({"gene_symbol": ["G1"], "big_col": [long_val], "fdr_q": [0.01]})
        parquet_path = tmp_path / "big.parquet"
        df.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["xlsx"],
        )
        result_df = pd.read_excel(paths[0])
        cell_val = str(result_df["big_col"].iloc[0])
        assert len(cell_val) <= _EXCEL_MAX_CELL_CHARS
        assert "TRUNCATED" in cell_val
        assert "50000" in cell_val

    def test_xlsx_truncation_warning_logged(self, tmp_path, caplog):
        long_val = "y" * 40000
        df = pd.DataFrame({"gene_symbol": ["G1"], "big_col": [long_val], "fdr_q": [0.01]})
        parquet_path = tmp_path / "big.parquet"
        df.to_parquet(parquet_path)

        import logging

        with caplog.at_level(logging.WARNING):
            export_results(
                result_type="gene",
                results_path=parquet_path,
                output_dir=tmp_path / "out",
                study_name="TEST",
                formats=["xlsx"],
            )
        assert any("truncation" in r.message.lower() for r in caplog.records)
        assert any("big_col" in r.message for r in caplog.records)

    def test_xlsx_no_truncation_for_short_values(self, tmp_path, gene_results):
        parquet_path = tmp_path / "genes.parquet"
        gene_results.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["xlsx"],
        )
        df = pd.read_excel(paths[0])
        for col in df.columns:
            for val in df[col].dropna():
                assert "TRUNCATED" not in str(val)

    def test_csv_not_truncated(self, tmp_path):
        """CSV must preserve the full value regardless of length."""
        long_val = "z" * 50000
        df = pd.DataFrame({"gene_symbol": ["G1"], "big_col": [long_val], "fdr_q": [0.01]})
        parquet_path = tmp_path / "big.parquet"
        df.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["csv"],
        )
        csv_df = pd.read_csv(paths[0])
        assert len(csv_df["big_col"].iloc[0]) == 50000

    def test_xlsx_control_char_sanitized(self, tmp_path):
        """XLSX must not raise on XML-illegal control characters."""
        dirty = "compound\x01TFA"
        df = pd.DataFrame({"drug_name": [dirty], "magma_fdr_q": [0.01]})
        parquet_path = tmp_path / "dirty.parquet"
        df.to_parquet(parquet_path)

        paths = export_results(
            result_type="drug",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["xlsx"],
        )
        assert len(paths) == 1
        assert paths[0].exists()
        result_df = pd.read_excel(paths[0])
        assert "\\x01" in result_df["drug_name"].iloc[0]

    def test_xlsx_sanitize_then_truncate(self, tmp_path):
        """Sanitization must occur before truncation."""
        val = "A\x01" + "B" * 50000
        df = pd.DataFrame({"gene_symbol": ["G1"], "big_col": [val], "fdr_q": [0.01]})
        parquet_path = tmp_path / "st.parquet"
        df.to_parquet(parquet_path)

        paths = export_results(
            result_type="gene",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["xlsx"],
        )
        result_df = pd.read_excel(paths[0])
        cell_val = str(result_df["big_col"].iloc[0])
        assert len(cell_val) <= _EXCEL_MAX_CELL_CHARS
        assert "\\x01" in cell_val
        assert "TRUNCATED" in cell_val

    def test_xlsx_sanitization_warning_logged(self, tmp_path, caplog):
        dirty = "name\x02here"
        df = pd.DataFrame({"drug_name": [dirty], "magma_fdr_q": [0.01]})
        parquet_path = tmp_path / "dirty.parquet"
        df.to_parquet(parquet_path)

        import logging

        with caplog.at_level(logging.WARNING):
            export_results(
                result_type="drug",
                results_path=parquet_path,
                output_dir=tmp_path / "out",
                study_name="TEST",
                formats=["xlsx"],
            )
        assert any("sanitization" in r.message.lower() for r in caplog.records)
        assert any("drug_name" in r.message for r in caplog.records)

    def test_csv_preserves_control_chars(self, tmp_path):
        """CSV must preserve control characters verbatim."""
        dirty = "compound\x01TFA"
        df = pd.DataFrame({"drug_name": [dirty], "magma_fdr_q": [0.01]})
        parquet_path = tmp_path / "dirty.parquet"
        df.to_parquet(parquet_path)

        paths = export_results(
            result_type="drug",
            results_path=parquet_path,
            output_dir=tmp_path / "out",
            study_name="TEST",
            formats=["csv"],
        )
        csv_df = pd.read_csv(paths[0])
        assert csv_df["drug_name"].iloc[0] == dirty

    def test_sanitize_helper_deterministic(self):
        assert _sanitize_excel_str("a\x00b\x01c") == "a\\x00b\\x01c"
        assert _sanitize_excel_str("no_special") == "no_special"
        assert _sanitize_excel_str("\x1f") == "\\x1f"


# ===================================================================
# TestExportCombined
# ===================================================================


class TestExportCombined:
    """Combined export tests."""

    def _make_combined(self, all_fixtures, study="TEST"):
        return CombinedResults(
            study_name=study,
            branches_present=["magma", "mr", "neg_correlation"],
            gene_results=all_fixtures["gene"],
            pathway_results=all_fixtures["pathway"],
            drug_enrichment=all_fixtures["drug"],
            atc_enrichment=all_fixtures["atc"],
            spredixcan_meta=all_fixtures["spredixcan"],
            neg_correlation_summary=all_fixtures["correlation"],
            mr_results=all_fixtures["mr"],
            mr_drug_matches=all_fixtures["mr_drugs"],
        )

    def test_xlsx_summary_sheet_first(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        paths = export_combined(combined, tmp_path, formats=["xlsx"])

        from openpyxl import load_workbook

        wb = load_workbook(paths[0])
        assert wb.sheetnames[0] == "Summary"

    def test_xlsx_all_branches(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        paths = export_combined(combined, tmp_path, formats=["xlsx"])

        from openpyxl import load_workbook

        wb = load_workbook(paths[0])
        assert len(wb.sheetnames) == 9

    def test_xlsx_partial_branches(self, tmp_path, gene_results, pathway_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            pathway_results=pathway_results,
        )
        paths = export_combined(combined, tmp_path, formats=["xlsx"])

        from openpyxl import load_workbook

        wb = load_workbook(paths[0])
        assert "Summary" in wb.sheetnames
        assert "Gene Results" in wb.sheetnames
        assert "Pathway Results" in wb.sheetnames
        assert "MR Results" not in wb.sheetnames

    def test_xlsx_frozen_panes(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        paths = export_combined(combined, tmp_path, formats=["xlsx"])

        from openpyxl import load_workbook

        wb = load_workbook(paths[0])
        for ws in wb.worksheets:
            assert ws.freeze_panes == "A2"

    def test_xlsx_summary_columns(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        paths = export_combined(combined, tmp_path, formats=["xlsx"])

        from openpyxl import load_workbook

        wb = load_workbook(paths[0])
        ws = wb["Summary"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, 11)]
        expected = [
            "drug_name",
            "drug_chembl_id",
            "branches_found_in",
            "magma_fdr_q",
            "neg_corr_best_rho",
            "neg_corr_best_tissue",
            "neg_corr_n_tissues_sig",
            "mr_confidence_tier",
            "max_phase",
            "atc_codes",
        ]
        assert headers == expected

    def test_combined_json_schema(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        paths = export_combined(combined, tmp_path, formats=["json"])

        json_path = [p for p in paths if p.suffix == ".json"][0]
        data = json.loads(json_path.read_text())
        assert data["result_type"] == "combined"
        assert data["study"] == "TEST"
        assert "branches_present" in data
        assert isinstance(data["data"], dict)
        assert "gene" in data["data"]

    def test_combined_csv_per_type(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        paths = export_combined(combined, tmp_path, formats=["csv"])

        csv_paths = [p for p in paths if p.suffix == ".csv"]
        assert len(csv_paths) == 8

    def _verdicts_df(self) -> pd.DataFrame:
        return pd.DataFrame({
            "gene_ensembl_id": ["ENSG001"],
            "gene_symbol": ["GENE1"],
            "eqtl_source": ["eqtlgen"],
            "mr_beta": [0.5],
            "mr_pval": [1e-10],
            "confidence_tier": ["high"],
            "verdict_status": ["actionable"],
            "n_drug_records_raw": [4],
            "n_filtered_matches": [2],
            "n_direction_inferable": [2],
            "n_direction_concordant": [1],
        })

    def test_verdicts_surfaced_in_csv_and_xlsx(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        combined.mr_target_verdicts = self._verdicts_df()
        paths = export_combined(combined, tmp_path, formats=["csv", "xlsx"])

        csv_names = {p.name for p in paths if p.suffix == ".csv"}
        assert any("mr_verdicts" in n for n in csv_names)

        from openpyxl import load_workbook

        wb = load_workbook([p for p in paths if p.suffix == ".xlsx"][0])
        assert "MR Target Verdicts" in wb.sheetnames

    def test_verdicts_discovered_by_combine_results(self, tmp_path):
        study = "TEST"
        mr_info = RESULT_FILE_MAP["mr_verdicts"]
        subdir = tmp_path / study / mr_info["subdir"]
        subdir.mkdir(parents=True)
        self._verdicts_df().to_parquet(
            subdir / TEST_FILENAMES["mr_verdicts"].format(study=study)
        )
        combined = combine_results(tmp_path, study)
        assert combined.mr_target_verdicts is not None
        assert combined.mr_target_verdicts.iloc[0]["verdict_status"] == "actionable"


# ===================================================================
# TestCombineResults
# ===================================================================


class TestCombineResults:
    """combine_results function tests."""

    def test_combine_all_branches(self, tmp_path, all_fixtures):
        study = "TEST_STUDY"
        _write_test_results(
            tmp_path,
            study,
            ["gene", "pathway", "drug", "atc", "spredixcan", "correlation", "mr", "mr_drugs"],
            all_fixtures,
        )
        combined = combine_results(tmp_path, study)

        assert "magma" in combined.branches_present
        assert "neg_correlation" in combined.branches_present
        assert "mr" in combined.branches_present
        assert combined.gene_results is not None
        assert combined.pathway_results is not None
        assert combined.drug_enrichment is not None
        assert combined.mr_results is not None

    def test_combine_magma_only(self, tmp_path, all_fixtures):
        study = "MAGMA_ONLY"
        _write_test_results(tmp_path, study, ["gene", "pathway", "drug", "atc"], all_fixtures)
        combined = combine_results(tmp_path, study)

        assert combined.branches_present == ["magma"]
        assert combined.gene_results is not None
        assert combined.neg_correlation_summary is None
        assert combined.mr_results is None

    def test_combine_neg_corr_only(self, tmp_path, all_fixtures):
        study = "NEG_CORR"
        _write_test_results(tmp_path, study, ["spredixcan", "correlation"], all_fixtures)
        combined = combine_results(tmp_path, study)

        assert combined.branches_present == ["neg_correlation"]
        assert combined.spredixcan_meta is not None
        assert combined.neg_correlation_summary is not None
        assert combined.gene_results is None

    def test_combine_mr_only(self, tmp_path, all_fixtures):
        study = "MR_ONLY"
        _write_test_results(tmp_path, study, ["mr", "mr_drugs"], all_fixtures)
        combined = combine_results(tmp_path, study)

        assert combined.branches_present == ["mr"]
        assert combined.mr_results is not None

    def test_combine_no_results_raises(self, tmp_path):
        study = "EMPTY"
        (tmp_path / study).mkdir()
        with pytest.raises(ValueError, match="No results found"):
            combine_results(tmp_path, study)

    def test_combine_drug_overlap_two_branches(self, tmp_path, all_fixtures):
        study = "OVERLAP"
        _write_test_results(tmp_path, study, ["drug", "mr_drugs"], all_fixtures)
        combined = combine_results(tmp_path, study)

        assert combined.drug_overlap is not None
        assert "in_magma" in combined.drug_overlap.columns
        assert "in_mr" in combined.drug_overlap.columns
        assert "n_branches" in combined.drug_overlap.columns
        overlap_drugs = combined.drug_overlap[combined.drug_overlap["n_branches"] >= 2]
        assert len(overlap_drugs) >= 1

    def test_combine_drug_overlap_no_overlap(self, tmp_path):
        study = "NO_OVERLAP"
        drug_df = pd.DataFrame({
            "drug_name": ["UniqueA"],
            "drug_chembl_id": ["CHEMBL999"],
            "magma_fdr_q": [0.001],
        })
        mr_df = pd.DataFrame({
            "drug_name": ["UniqueB"],
            "drug_chembl_id": ["CHEMBL888"],
            "confidence_tier": ["high"],
        })
        drug_info = RESULT_FILE_MAP["drug"]
        mr_info = RESULT_FILE_MAP["mr_drugs"]
        (tmp_path / study / drug_info["subdir"]).mkdir(parents=True)
        (tmp_path / study / mr_info["subdir"]).mkdir(parents=True)
        drug_df.to_parquet(
            tmp_path / study / drug_info["subdir"] / f"{study}_drug_enrichment.parquet"
        )
        mr_df.to_parquet(
            tmp_path / study / mr_info["subdir"] / f"{study}_mr_drug_matches.parquet"
        )
        combined = combine_results(tmp_path, study)

        assert combined.drug_overlap is not None
        assert len(combined.drug_overlap) == 0

    def test_combine_drug_overlap_single_branch(self, tmp_path, all_fixtures):
        study = "SINGLE"
        _write_test_results(tmp_path, study, ["drug"], all_fixtures)
        combined = combine_results(tmp_path, study)

        assert combined.drug_overlap is None

    def test_combine_metadata_recorded(self, tmp_path, all_fixtures):
        study = "META"
        _write_test_results(tmp_path, study, ["gene", "mr"], all_fixtures)
        combined = combine_results(tmp_path, study)

        assert "study_name" in combined.metadata
        assert "branches_present" in combined.metadata
        assert "timestamp" in combined.metadata
        assert "files_loaded" in combined.metadata

    def test_combine_corrupt_file_skipped(self, tmp_path, all_fixtures):
        study = "CORRUPT"
        _write_test_results(tmp_path, study, ["gene"], all_fixtures)
        corrupt_info = RESULT_FILE_MAP["mr"]
        corrupt_dir = tmp_path / study / corrupt_info["subdir"]
        corrupt_dir.mkdir(parents=True, exist_ok=True)
        corrupt_file = corrupt_dir / f"{study}_mr_results.parquet"
        corrupt_file.write_bytes(b"not a parquet file")

        combined = combine_results(tmp_path, study)
        assert combined.gene_results is not None
        assert combined.mr_results is None


# ===================================================================
# TestHtmlReport
# ===================================================================


class TestHtmlReport:
    """HTML report generation tests."""

    def _make_combined(self, all_fixtures, study="TEST"):
        return CombinedResults(
            study_name=study,
            branches_present=["magma", "mr", "neg_correlation"],
            gene_results=all_fixtures["gene"],
            pathway_results=all_fixtures["pathway"],
            drug_enrichment=all_fixtures["drug"],
            atc_enrichment=all_fixtures["atc"],
            spredixcan_meta=all_fixtures["spredixcan"],
            neg_correlation_summary=all_fixtures["correlation"],
            mr_results=all_fixtures["mr"],
            mr_drug_matches=all_fixtures["mr_drugs"],
            metadata={"study_name": study, "branches_present": ["magma", "mr", "neg_correlation"]},
        )

    def test_html_full_report(self, tmp_path, all_fixtures):
        combined = self._make_combined(all_fixtures)
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        assert out.exists()
        content = out.read_text()
        assert "magma-genes" in content
        assert "pathways" in content
        assert "drug-enrichment" in content
        assert "atc-enrichment" in content
        assert "spredixcan" in content
        assert "neg-correlation" in content
        assert "mendelian-randomisation" in content

    def test_html_partial_report_magma_only(self, tmp_path, gene_results, pathway_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            pathway_results=pathway_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert "magma-genes" in content
        assert "mendelian-randomisation" not in content
        assert "neg-correlation" not in content

    def test_html_partial_report_neg_corr_only(
        self, tmp_path, spredixcan_meta, neg_correlation_summary
    ):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["neg_correlation"],
            spredixcan_meta=spredixcan_meta,
            neg_correlation_summary=neg_correlation_summary,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert "spredixcan" in content
        assert "neg-correlation" in content
        assert "magma-genes" not in content

    def test_html_single_branch(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert "convergence" not in content

    def test_html_valid_markup(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()

        class HTMLValidator(html.parser.HTMLParser):
            def __init__(self):
                super().__init__()
                self.errors = []

            def handle_starttag(self, tag, attrs):
                pass

            def handle_endtag(self, tag):
                pass

        validator = HTMLValidator()
        validator.feed(content)
        assert not validator.errors

    def test_html_figure_embedding(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()
        _create_dummy_png(plot_dir / "manhattan.png")

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert "data:image/png;base64," in content

    def test_html_missing_figures(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert '<img src="data:image/png' not in content

    def test_html_table_formatting(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert "3.40e-06" in content or "3.4e-06" in content

    def test_html_self_contained(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert 'href="http' not in content
        assert 'src="http' not in content
        assert '<link rel="stylesheet"' not in content


# ===================================================================
# Helper function tests
# ===================================================================


class TestHelpers:
    """Tests for internal helper functions."""

    def test_convert_numpy_types_int(self):
        result = _convert_numpy_types(np.int64(42))
        assert result == 42
        assert isinstance(result, int)

    def test_convert_numpy_types_float(self):
        result = _convert_numpy_types(np.float64(3.14))
        assert result == pytest.approx(3.14)
        assert isinstance(result, float)

    def test_convert_numpy_types_nan(self):
        result = _convert_numpy_types(np.float64("nan"))
        assert result is None

    def test_convert_numpy_types_bool(self):
        result = _convert_numpy_types(np.bool_(True))
        assert result is True
        assert isinstance(result, bool)

    def test_convert_numpy_types_nested(self):
        data = [{"a": np.int64(1), "b": [np.float64(2.0), np.bool_(False)]}]
        result = _convert_numpy_types(data)
        assert result == [{"a": 1, "b": [2.0, False]}]

    def test_build_summary_gene(self, gene_results):
        summary = _build_summary(gene_results, "gene")
        assert summary["n_genes_tested"] == 5
        assert summary["n_significant"] == 4

    def test_build_summary_mr(self, mr_results):
        summary = _build_summary(mr_results, "mr")
        assert summary["n_genes_tested"] == 3
        assert summary["n_significant"] == 2
        assert summary["n_colocalised"] == 2

    def test_build_summary_missing_column(self):
        df = pd.DataFrame({"x": [1, 2, 3]})
        summary = _build_summary(df, "gene")
        assert summary["n_significant"] == 0

    def test_dataframe_to_records(self, gene_results):
        records = _dataframe_to_records(gene_results)
        assert len(records) == 5
        assert all(isinstance(r, dict) for r in records)
        assert all(isinstance(r["magma_z"], (int, float)) for r in records)

    def test_embed_figure_exists(self, tmp_path):
        _create_dummy_png(tmp_path / "test.png")
        result = _embed_figure(tmp_path, "test.png")
        assert result is not None
        assert result.startswith("data:image/png;base64,")

    def test_embed_figure_missing(self, tmp_path):
        result = _embed_figure(tmp_path, "nonexistent.png")
        assert result is None

    def test_render_table(self, gene_results):
        html = _render_table(gene_results)
        assert "<table" in html
        assert "<thead>" in html
        assert "<tbody>" in html
        assert "GENE1" in html

    def test_render_table_max_rows(self, gene_results):
        html = _render_table(gene_results, max_rows=2)
        assert "GENE3" not in html

    def test_render_table_selected_columns(self, gene_results):
        html = _render_table(gene_results, columns=["gene_symbol", "magma_p"])
        assert "biotype" not in html.lower().replace("biotype", "")
        assert "Gene Symbol" in html

    def test_render_table_sort_by(self, gene_results):
        html = _render_table(
            gene_results,
            columns=["gene_symbol", "magma_p"],
            sort_by="magma_p",
            ascending=True,
        )
        rows = html.split("<tr>")
        gene_rows = [r for r in rows if "GENE" in r]
        assert "GENE4" in gene_rows[0]

    def test_render_table_sort_descending(self, gene_results):
        html = _render_table(
            gene_results,
            columns=["gene_symbol", "magma_p"],
            sort_by="magma_p",
            ascending=False,
        )
        rows = html.split("<tr>")
        gene_rows = [r for r in rows if "GENE" in r]
        assert "GENE5" in gene_rows[0]

    def test_render_table_max_rows_zero_renders_all(self, gene_results):
        html = _render_table(gene_results, max_rows=0)
        for name in ["GENE1", "GENE2", "GENE3", "GENE4", "GENE5"]:
            assert name in html


# ===================================================================
# Fix-A / Fix-B regression tests
# ===================================================================


class TestHtmlTableEscaping:
    """Fix A: tables must render as raw HTML, not escaped text."""

    def test_tables_contain_raw_html_tags(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert '<table class="data-table">' in content
        assert "&lt;table" not in content

    def test_all_sections_have_unescaped_tables(self, tmp_path, all_fixtures):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma", "mr", "neg_correlation"],
            gene_results=all_fixtures["gene"],
            pathway_results=all_fixtures["pathway"],
            drug_enrichment=all_fixtures["drug"],
            atc_enrichment=all_fixtures["atc"],
            spredixcan_meta=all_fixtures["spredixcan"],
            neg_correlation_summary=all_fixtures["correlation"],
            mr_results=all_fixtures["mr"],
            mr_drug_matches=all_fixtures["mr_drugs"],
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        assert content.count('<table class="data-table">') >= 8
        assert "&lt;table" not in content
        assert "&lt;thead" not in content
        assert "&lt;td" not in content

    def test_data_markup_injection_escaped(self, tmp_path):
        """Regression: data values containing HTML must be escaped in cells."""
        injected = pd.DataFrame(
            {
                "gene_symbol": ["<b>INJECT</b>", "BRCA1", '<script>alert("xss")</script>'],
                "chr": [1, 2, 3],
                "start": [100, 200, 300],
                "magma_z": [5.0, 3.0, 1.0],
                "magma_p": [1e-7, 1e-4, 0.05],
                "fdr_q": [1e-6, 1e-3, 0.1],
                "n_snps": [10, 20, 30],
                "biotype": ["protein_coding"] * 3,
                "in_mhc": [False] * 3,
                "annotation_mode": ["proximity"] * 3,
            }
        )
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=injected,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()

        assert "&lt;b&gt;INJECT&lt;/b&gt;" in content
        assert "<b>INJECT</b>" not in content

        assert "&lt;script&gt;" in content
        assert "<script>alert" not in content

        assert '<table class="data-table">' in content


class TestHtmlTableSorting:
    """Fix B: tables must be sorted by the spec-defined metrics."""

    def test_gene_table_sorted_by_magma_p(self, tmp_path, gene_results):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma"],
            gene_results=gene_results,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        pos_gene4 = content.index("GENE4")
        pos_gene1 = content.index("GENE1")
        pos_gene5 = content.index("GENE5")
        assert pos_gene4 < pos_gene1 < pos_gene5

    def test_neg_corr_sorted_most_negative_first(
        self, tmp_path, neg_correlation_summary
    ):
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["neg_correlation"],
            neg_correlation_summary=neg_correlation_summary,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        pos_a = content.index("DrugA")
        pos_x = content.index("DrugX")
        assert pos_a < pos_x

    def test_convergence_table_not_truncated(self, tmp_path):
        drug_df = pd.DataFrame(
            {
                "drug_name": [f"Drug{i}" for i in range(30)],
                "drug_chembl_id": [f"CHEMBL{i}" for i in range(30)],
                "in_magma": [True] * 30,
                "in_neg_corr": [True] * 30,
                "in_mr": [False] * 30,
                "n_branches": [2] * 30,
                "magma_fdr_q": [0.01] * 30,
                "neg_corr_best_rho": [-0.3] * 30,
                "mr_confidence_tier": [None] * 30,
            }
        )
        combined = CombinedResults(
            study_name="TEST",
            branches_present=["magma", "neg_correlation"],
            drug_overlap=drug_df,
            metadata={},
        )
        plot_dir = tmp_path / "plots"
        plot_dir.mkdir()

        out = generate_html_report(combined, plot_dir, tmp_path / "report.html")
        content = out.read_text()
        for i in range(30):
            assert f"Drug{i}" in content
