"""RepoGen reporting - export results to CSV, JSON, and XLSX."""

from __future__ import annotations

import argparse
import json
import logging
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import TYPE_CHECKING, Any

import numpy as np
import pandas as pd

from repogen.utils.logging import setup_logging

if TYPE_CHECKING:
    from repogen.reporting.combine_results import CombinedResults

logger = setup_logging(__name__)


def _filter_to_headline(df: pd.DataFrame | None) -> pd.DataFrame | None:
    """Filter the drug-enrichment dataframe to headline rows.

    Idempotent and graceful:
    - Returns ``df`` unchanged when ``df`` is None.
    - Returns ``df`` unchanged when the ``passes_headline_min_genes``
      column is absent (e.g. legacy parquet files written before the
      headline threshold existed).  This keeps backwards compatibility with archived
      results.
    - Otherwise returns a copy filtered to rows where
      ``passes_headline_min_genes`` is truthy (NaN treated as False).
    """
    if df is None or "passes_headline_min_genes" not in df.columns:
        return df
    mask = df["passes_headline_min_genes"].fillna(False).astype(bool)
    return df.loc[mask].copy()


VALID_RESULT_TYPES = (
    "gene",
    "pathway",
    "drug",
    "atc",
    "correlation",
    "correlation_per_tissue",
    "spredixcan",
    "spredixcan_per_tissue",
    "mr",
    "mr_drugs",
)

_SIGNIFICANCE_CONFIG: dict[str, dict[str, Any]] = {
    "gene": {
        "summary_fields": {"count": "n_genes_tested", "sig": "n_significant"},
        "sig_col": "fdr_q",
        "sig_op": "lt",
        "sig_val": 0.05,
    },
    "pathway": {
        "summary_fields": {"count": "n_pathways_tested", "sig": "n_significant"},
        "sig_col": "fdr_q",
        "sig_op": "lt",
        "sig_val": 0.05,
    },
    "drug": {
        "summary_fields": {"count": "n_drugs_tested", "sig": "n_significant"},
        "sig_col": "magma_fdr_q",
        "sig_op": "lt",
        "sig_val": 0.05,
    },
    "atc": {
        "summary_fields": {"count": "n_classes_tested", "sig": "n_significant"},
        "sig_col": "gls_fdr",
        "sig_op": "lt",
        "sig_val": 0.05,
    },
    "correlation": {
        "summary_fields": {"count": "n_drugs_tested", "sig": "n_significant"},
        "sig_col": "n_tissues_fdr_significant",
        "sig_op": "gt",
        "sig_val": 0,
    },
    "correlation_per_tissue": {
        "summary_fields": {"count": "n_drug_tissue_pairs", "sig": "n_significant"},
        "sig_col": "fdr_global",
        "sig_op": "lt",
        "sig_val": 0.05,
    },
    "spredixcan": {
        "summary_fields": {"count": "n_genes", "sig": "n_significant"},
        "sig_col": "meta_pvalue",
        "sig_op": "lt",
        "sig_val": 0.05,
    },
    "spredixcan_per_tissue": {
        "summary_fields": {"count": "n_gene_tissue_pairs", "sig": "n_significant"},
        "sig_col": "pvalue",
        "sig_op": "lt",
        "sig_val": 0.05,
    },
    "mr": {
        "summary_fields": {
            "count": "n_genes_tested",
            "sig": "n_significant",
            "extra": {"n_colocalised": ("pp_h4", "gte", 0.8)},
        },
        "sig_col": "mr_significant",
        "sig_op": "eq",
        "sig_val": True,
    },
    "mr_drugs": {
        "summary_fields": {
            "count": "n_drugs_matched",
            "sig": "n_high_confidence",
        },
        "sig_col": "confidence_tier",
        "sig_op": "eq",
        "sig_val": "high",
    },
}


def export_results(
    result_type: str,
    results_path: Path,
    metadata_path: Path | None = None,
    output_dir: Path = Path("."),
    study_name: str = "study",
    formats: list[str] | None = None,
) -> list[Path]:
    """Export a single result type to CSV, JSON, and/or XLSX.

    Args:
        result_type: One of the VALID_RESULT_TYPES identifiers.
        results_path: Path to the Parquet file from Phase 3.
        metadata_path: Path to the JSON metadata sidecar (optional).
        output_dir: Directory to write output files.
        study_name: Study identifier for file naming.
        formats: List of output formats (``"csv"``, ``"json"``, ``"xlsx"``).
            Default: ``["csv", "json"]``.

    Returns:
        List of paths to created files.

    Raises:
        FileNotFoundError: If *results_path* does not exist.
        ValueError: If *result_type* is not recognised.
    """
    if result_type not in VALID_RESULT_TYPES:
        raise ValueError(
            f"Unknown result_type '{result_type}'. "
            f"Must be one of: {', '.join(VALID_RESULT_TYPES)}"
        )

    results_path = Path(results_path)
    if not results_path.is_file():
        raise FileNotFoundError(f"Results file not found: {results_path}")

    if formats is None:
        formats = ["csv", "json"]

    df = pd.read_parquet(results_path)

    if result_type == "drug":
        n_pre = len(df)
        df = _filter_to_headline(df)
        if df is not None and len(df) < n_pre:
            logger.info(
                "filtered %d sub-headline drugs from "
                "drug export (kept %d headline drugs).",
                n_pre - len(df), len(df),
            )

    metadata: dict = {}
    if metadata_path is not None:
        metadata_path = Path(metadata_path)
        if metadata_path.is_file():
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    created: list[Path] = []
    for fmt in formats:
        if fmt == "csv":
            created.append(_export_csv(df, result_type, output_dir, study_name))
        elif fmt == "json":
            created.append(
                _export_json(df, result_type, metadata, output_dir, study_name)
            )
        elif fmt == "xlsx":
            created.append(
                _export_single_xlsx(df, result_type, output_dir, study_name)
            )
        else:
            logger.warning("Unsupported export format '%s', skipping.", fmt)

    logger.info(
        "Exported %s (%d rows) to %d file(s).", result_type, len(df), len(created)
    )
    return created


def _export_csv(
    df: pd.DataFrame,
    result_type: str,
    output_dir: Path,
    study_name: str,
) -> Path:
    """Write one CSV file. List-like columns are JSON-stringified."""
    output_path = output_dir / f"{study_name}_{result_type}.csv"

    df_csv = df.copy()
    for col in df_csv.columns:
        col_dtype = str(df_csv[col].dtype)
        needs_stringify = False
        if col_dtype == "object" or col_dtype.startswith("list") or "list" in col_dtype:
            non_null = df_csv[col].dropna()
            if not non_null.empty:
                sample = non_null.iloc[0]
                if isinstance(sample, (list, dict, np.ndarray)):
                    needs_stringify = True
        if needs_stringify:
            df_csv[col] = df_csv[col].apply(
                lambda x: json.dumps(x if not isinstance(x, np.ndarray) else x.tolist())
                if isinstance(x, (list, dict, np.ndarray))
                else x
            )

    df_csv.to_csv(output_path, index=False)
    return output_path


def _export_json(
    df: pd.DataFrame,
    result_type: str,
    metadata: dict,
    output_dir: Path,
    study_name: str,
) -> Path:
    """Write structured JSON following the Section 3.2G schema."""
    output_path = output_dir / f"{study_name}_{result_type}.json"

    output = {
        "result_type": result_type,
        "study": study_name,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "parameters": metadata.get("parameters", metadata.get("config", {})),
        "summary": _build_summary(df, result_type),
        "data": _dataframe_to_records(df),
    }

    output_path.write_text(
        json.dumps(output, indent=2, default=str, ensure_ascii=False),
        encoding="utf-8",
    )
    return output_path


def _export_single_xlsx(
    df: pd.DataFrame,
    result_type: str,
    output_dir: Path,
    study_name: str,
) -> Path:
    """Write a single-sheet XLSX workbook for one result type.

    Uses the shared ``_write_sheet`` helper so formatting and Excel
    cell-length truncation are identical to the combined XLSX workbook.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output_path = output_dir / f"{study_name}_{result_type}.xlsx"
    wb = Workbook()

    header_font = Font(bold=True)
    header_fill = PatternFill(
        start_color="D5E8F0", end_color="D5E8F0", fill_type="solid"
    )
    header_alignment = Alignment(horizontal="center", wrap_text=True)

    ws = wb.active
    ws.title = result_type
    _write_sheet(ws, df, header_font, header_fill, header_alignment)

    wb.save(output_path)
    logger.info("Per-result XLSX written to %s", output_path)
    return output_path


def _build_summary(df: pd.DataFrame, result_type: str) -> dict:
    """Compute summary statistics appropriate to each result_type."""
    config = _SIGNIFICANCE_CONFIG.get(result_type, {})
    if not config:
        return {"n_rows": len(df)}

    fields = config["summary_fields"]
    sig_col = config["sig_col"]
    sig_op = config["sig_op"]
    sig_val = config["sig_val"]

    summary: dict[str, Any] = {fields["count"]: len(df)}

    if sig_col in df.columns:
        if sig_op == "lt":
            n_sig = int((df[sig_col] < sig_val).sum())
        elif sig_op == "gt":
            n_sig = int((df[sig_col] > sig_val).sum())
        elif sig_op == "eq":
            n_sig = int((df[sig_col] == sig_val).sum())
        elif sig_op == "gte":
            n_sig = int((df[sig_col] >= sig_val).sum())
        else:
            n_sig = 0
        summary[fields["sig"]] = n_sig
    else:
        logger.warning(
            "Significance column '%s' not found in %s results; "
            "setting %s to 0.",
            sig_col,
            result_type,
            fields["sig"],
        )
        summary[fields["sig"]] = 0

    extra = fields.get("extra")
    if isinstance(extra, dict):
        for key, (col, op, val) in extra.items():
            if col in df.columns:
                if op == "gte":
                    summary[key] = int((df[col] >= val).sum())
                elif op == "lt":
                    summary[key] = int((df[col] < val).sum())
            else:
                summary[key] = 0

    return summary


def _dataframe_to_records(df: pd.DataFrame) -> list[dict]:
    """Convert DataFrame to JSON-serialisable list of dicts."""
    records = df.where(df.notna(), None).to_dict(orient="records")
    return _convert_numpy_types(records)


def _convert_numpy_types(obj: Any) -> Any:
    """Recursively convert numpy types to Python native types."""
    if isinstance(obj, list):
        return [_convert_numpy_types(item) for item in obj]
    if isinstance(obj, dict):
        return {k: _convert_numpy_types(v) for k, v in obj.items()}
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating,)):
        if np.isnan(obj):
            return None
        return float(obj)
    if isinstance(obj, (np.bool_,)):
        return bool(obj)
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    return obj


# ---------------------------------------------------------------------------
# Combined export
# ---------------------------------------------------------------------------


def export_combined(
    combined: CombinedResults,
    output_dir: Path,
    formats: list[str] | None = None,
) -> list[Path]:
    """Export aggregated results from combine_results output.

    Produces multi-sheet XLSX workbook + combined JSON.

    Args:
        combined: CombinedResults dataclass from combine_results.py.
        output_dir: Directory to write output files.
        formats: List of output formats. Default: ``["csv", "json", "xlsx"]``.

    Returns:
        List of paths to created files.

    Raises:
        ValueError: If combined has no branches_present (nothing to export).
    """
    if not combined.branches_present:
        raise ValueError("CombinedResults has no branches_present - nothing to export.")

    if formats is None:
        formats = ["csv", "json", "xlsx"]

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    created: list[Path] = []

    _FIELD_TYPE_MAP = {
        "gene_results": "gene",
        "pathway_results": "pathway",
        "drug_enrichment": "drug",
        "atc_enrichment": "atc",
        "spredixcan_meta": "spredixcan",
        "neg_correlation_summary": "correlation",
        "mr_results": "mr",
        "mr_drug_matches": "mr_drugs",
        "mr_target_verdicts": "mr_verdicts",
    }

    if "csv" in formats:
        for attr, rtype in _FIELD_TYPE_MAP.items():
            df = getattr(combined, attr, None)
            if df is not None:
                if attr == "drug_enrichment":
                    df = _filter_to_headline(df)
                path = _export_csv(
                    df, f"combined_{rtype}", output_dir, combined.study_name
                )
                created.append(path)

    if "json" in formats:
        combined_json = _build_combined_json(combined, _FIELD_TYPE_MAP)
        json_path = output_dir / f"{combined.study_name}_combined.json"
        json_path.write_text(
            json.dumps(combined_json, indent=2, default=str, ensure_ascii=False),
            encoding="utf-8",
        )
        created.append(json_path)

    if "xlsx" in formats:
        created.append(_export_xlsx(combined, output_dir))

    logger.info(
        "Combined export for '%s' produced %d file(s).",
        combined.study_name,
        len(created),
    )
    return created


def _build_combined_json(
    combined: CombinedResults,
    field_type_map: dict[str, str],
) -> dict:
    """Build combined JSON structure."""
    data_dict: dict[str, Any] = {}
    for attr, rtype in field_type_map.items():
        df = getattr(combined, attr, None)
        if df is not None:
            if attr == "drug_enrichment":
                df = _filter_to_headline(df)
            data_dict[rtype] = {
                "summary": _build_summary(df, rtype),
                "data": _dataframe_to_records(df),
            }

    return {
        "result_type": "combined",
        "study": combined.study_name,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "branches_present": combined.branches_present,
        "parameters": combined.metadata.get("config", {}),
        "data": data_dict,
    }


def _export_xlsx(
    combined: CombinedResults,
    output_dir: Path,
) -> Path:
    """Write multi-sheet XLSX workbook with openpyxl.

    Summary sheet FIRST - this is the actionable output researchers want.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    output_path = output_dir / f"{combined.study_name}_combined.xlsx"
    wb = Workbook()

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D5E8F0", end_color="D5E8F0", fill_type="solid")
    header_alignment = Alignment(horizontal="center", wrap_text=True)

    summary_df = _build_summary_sheet(combined)
    ws_summary = wb.active
    ws_summary.title = "Summary"
    _write_sheet(ws_summary, summary_df, header_font, header_fill, header_alignment)

    sheets = [
        ("Gene Results", combined.gene_results),
        ("Pathway Results", combined.pathway_results),
        ("Drug Enrichment", _filter_to_headline(combined.drug_enrichment)),
        ("ATC Enrichment", combined.atc_enrichment),
        ("S-PrediXcan Meta", combined.spredixcan_meta),
        ("Neg Correlation", combined.neg_correlation_summary),
        ("MR Results", combined.mr_results),
        ("MR Drug Matches", combined.mr_drug_matches),
        ("MR Target Verdicts", combined.mr_target_verdicts),
    ]

    for sheet_name, df in sheets:
        if df is not None:
            ws = wb.create_sheet(title=sheet_name)
            _write_sheet(ws, df, header_font, header_fill, header_alignment)

    wb.save(output_path)
    logger.info("XLSX workbook written to %s", output_path)
    return output_path


_EXCEL_MAX_CELL_CHARS = 32767
_EXCEL_ILLEGAL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _sanitize_excel_str(value: str) -> str:
    """Replace XML-illegal control characters with deterministic ``\\xNN`` escapes."""
    return _EXCEL_ILLEGAL_CHARS_RE.sub(lambda m: f"\\x{ord(m.group()):02x}", value)


def _write_sheet(
    ws: Any,
    df: pd.DataFrame,
    header_font: Any,
    header_fill: Any,
    header_alignment: Any,
) -> dict[str, int]:
    """Write a DataFrame to an openpyxl worksheet with formatting.

    Processing order for string cells:
    1. JSON-stringify list/dict/ndarray values.
    2. Sanitize XML-illegal control characters (``\\xNN`` escapes).
    3. Truncate to Excel 32 767-character cell limit.
    4. Write to worksheet.

    Returns:
        Dict mapping column names to the number of cells that were
        truncated to satisfy the Excel cell-length limit.
        Empty dict when no truncation occurred.
    """
    truncation_counts: dict[str, int] = {}
    sanitization_counts: dict[str, int] = {}
    col_names = list(df.columns)

    df_out = df.copy()
    for col in df_out.columns:
        col_dtype = str(df_out[col].dtype)
        needs_stringify = False
        if col_dtype == "object" or col_dtype.startswith("list") or "list" in col_dtype:
            non_null = df_out[col].dropna()
            if not non_null.empty:
                sample = non_null.iloc[0]
                if isinstance(sample, (list, dict, np.ndarray)):
                    needs_stringify = True
        if needs_stringify:
            df_out[col] = df_out[col].apply(
                lambda x: json.dumps(x if not isinstance(x, np.ndarray) else x.tolist())
                if isinstance(x, (list, dict, np.ndarray))
                else x
            )

    for col_idx, col_name in enumerate(df_out.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col_name))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, row in enumerate(df_out.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            cell_value = value
            if isinstance(cell_value, (np.integer,)):
                cell_value = int(cell_value)
            elif isinstance(cell_value, (np.floating,)):
                cell_value = None if np.isnan(cell_value) else float(cell_value)
            elif isinstance(cell_value, (np.bool_,)):
                cell_value = bool(cell_value)
            elif pd.isna(cell_value) if not isinstance(cell_value, str) else False:
                cell_value = None

            if isinstance(cell_value, str):
                if _EXCEL_ILLEGAL_CHARS_RE.search(cell_value):
                    cell_value = _sanitize_excel_str(cell_value)
                    col_name = col_names[col_idx - 1]
                    sanitization_counts[col_name] = (
                        sanitization_counts.get(col_name, 0) + 1
                    )

                if len(cell_value) > _EXCEL_MAX_CELL_CHARS:
                    original_len = len(cell_value)
                    suffix = f"...[TRUNCATED from {original_len} chars]"
                    cell_value = (
                        cell_value[: _EXCEL_MAX_CELL_CHARS - len(suffix)] + suffix
                    )
                    col_name = col_names[col_idx - 1]
                    truncation_counts[col_name] = (
                        truncation_counts.get(col_name, 0) + 1
                    )

            ws.cell(row=row_idx, column=col_idx, value=cell_value)

    from openpyxl.utils import get_column_letter

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    for col_idx, col_name in enumerate(df_out.columns, start=1):
        width = min(max(len(str(col_name)) + 4, 10), 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    sheet_label = ws.title or "sheet"

    if sanitization_counts:
        detail = ", ".join(f"{c}: {n}" for c, n in sanitization_counts.items())
        logger.warning(
            "Illegal XML control-char sanitization in '%s': %s",
            sheet_label,
            detail,
        )

    if truncation_counts:
        detail = ", ".join(f"{c}: {n}" for c, n in truncation_counts.items())
        logger.warning(
            "Excel cell-length truncation in '%s': %s (limit %d chars)",
            sheet_label,
            detail,
            _EXCEL_MAX_CELL_CHARS,
        )

    return truncation_counts


def _build_summary_sheet(combined: CombinedResults) -> pd.DataFrame:
    """Build the cross-branch drug overview for the Summary sheet."""
    records: list[dict[str, Any]] = []

    if combined.drug_enrichment is not None:
        de = _filter_to_headline(combined.drug_enrichment)
        sig_col = "magma_fdr_q"
        if de is not None and sig_col in de.columns:
            sig_drugs = de[de[sig_col] < 0.05]
        else:
            sig_drugs = (de if de is not None else combined.drug_enrichment).head(0)

        for row in sig_drugs.itertuples(index=False):
            name = getattr(row, "drug_name", None)
            chembl = getattr(row, "drug_chembl_id", None)
            atc = getattr(row, "atc_codes", None)
            if isinstance(atc, list):
                atc = ", ".join(str(c) for c in atc)
            elif isinstance(atc, str):
                try:
                    parsed = json.loads(atc)
                    if isinstance(parsed, list):
                        atc = ", ".join(str(c) for c in parsed)
                except (json.JSONDecodeError, TypeError):
                    pass

            records.append({
                "drug_name": name,
                "drug_chembl_id": chembl,
                "_branches": {"MAGMA"},
                "magma_fdr_q": getattr(row, sig_col, None),
                "neg_corr_best_rho": None,
                "neg_corr_best_tissue": None,
                "neg_corr_n_tissues_sig": None,
                "mr_confidence_tier": None,
                "max_phase": getattr(row, "max_phase", None),
                "atc_codes": atc,
            })

    if combined.neg_correlation_summary is not None:
        nc = combined.neg_correlation_summary
        sig_col_nc = "n_tissues_fdr_significant"
        if sig_col_nc in nc.columns:
            sig_nc = nc[nc[sig_col_nc] > 0]
        else:
            sig_nc = nc.head(0)

        for row in sig_nc.itertuples(index=False):
            name = getattr(row, "drug_name", None)
            chembl = getattr(row, "drug_chembl_id", None)
            records.append({
                "drug_name": name,
                "drug_chembl_id": chembl,
                "_branches": {"Neg Correlation"},
                "magma_fdr_q": None,
                "neg_corr_best_rho": getattr(row, "best_spearman_rho", None),
                "neg_corr_best_tissue": getattr(row, "best_tissue", None),
                "neg_corr_n_tissues_sig": getattr(row, sig_col_nc, None),
                "mr_confidence_tier": None,
                "max_phase": None,
                "atc_codes": None,
            })

    if combined.mr_drug_matches is not None:
        mr = combined.mr_drug_matches
        for row in mr.itertuples(index=False):
            name = getattr(row, "drug_name", None)
            chembl = getattr(row, "drug_chembl_id", None)
            records.append({
                "drug_name": name,
                "drug_chembl_id": chembl,
                "_branches": {"MR"},
                "magma_fdr_q": None,
                "neg_corr_best_rho": None,
                "neg_corr_best_tissue": None,
                "neg_corr_n_tissues_sig": None,
                "mr_confidence_tier": getattr(row, "confidence_tier", None),
                "max_phase": getattr(row, "max_phase", None),
                "atc_codes": None,
            })

    if not records:
        return pd.DataFrame({
            "drug_name": ["No drugs reached significance threshold in any branch."],
            "drug_chembl_id": [None],
            "branches_found_in": [None],
            "magma_fdr_q": [None],
            "neg_corr_best_rho": [None],
            "neg_corr_best_tissue": [None],
            "neg_corr_n_tissues_sig": [None],
            "mr_confidence_tier": [None],
            "max_phase": [None],
            "atc_codes": [None],
        })

    merged = _merge_drug_records(records)

    result_cols = [
        "drug_name",
        "drug_chembl_id",
        "branches_found_in",
        "magma_fdr_q",
        "neg_corr_best_rho",
        "neg_corr_best_tissue",
        "neg_corr_n_tissues_sig",
        "mr_confidence_tier",
        "max_phase",
        "atc_codes",
    ]
    df = pd.DataFrame(merged)

    for col in result_cols:
        if col not in df.columns:
            df[col] = None

    df = df[result_cols]

    df["_n_branches"] = df["branches_found_in"].apply(
        lambda x: len(x.split(", ")) if isinstance(x, str) else 0
    )
    df = df.sort_values(
        ["_n_branches", "magma_fdr_q"],
        ascending=[False, True],
        na_position="last",
    ).drop(columns=["_n_branches"])

    return df.reset_index(drop=True)


def _merge_drug_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Merge drug records by chembl_id primary, name fallback."""
    by_chembl: dict[str, dict] = {}
    by_name: dict[str, dict] = {}

    for rec in records:
        chembl = rec.get("drug_chembl_id")
        name = rec.get("drug_name")
        name_norm = name.strip().lower() if isinstance(name, str) else None
        branches = rec.pop("_branches", set())

        matched = None
        if chembl and chembl in by_chembl:
            matched = by_chembl[chembl]
        elif name_norm and name_norm in by_name:
            matched = by_name[name_norm]

        if matched is not None:
            matched["_branches"] |= branches
            for k, v in rec.items():
                if v is not None and matched.get(k) is None:
                    matched[k] = v
        else:
            rec["_branches"] = branches
            if chembl:
                by_chembl[chembl] = rec
            if name_norm:
                by_name[name_norm] = rec

    all_records = list({id(r): r for r in list(by_chembl.values()) + list(by_name.values())}.values())

    for rec in all_records:
        br = rec.pop("_branches", set())
        rec["branches_found_in"] = ", ".join(sorted(br))

    return all_records


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Export RepoGen analysis results to CSV/JSON/XLSX."
    )
    parser.add_argument(
        "--result-type",
        required=True,
        choices=VALID_RESULT_TYPES,
        help="Type of result to export.",
    )
    parser.add_argument("--results-path", required=True, type=Path)
    parser.add_argument("--metadata-path", type=Path, default=None)
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    parser.add_argument("--study-name", default="study")
    parser.add_argument(
        "--formats",
        nargs="+",
        default=["csv", "json"],
        help="Output formats (csv, json, xlsx). Default: csv json.",
    )

    args = parser.parse_args()
    paths = export_results(
        result_type=args.result_type,
        results_path=args.results_path,
        metadata_path=args.metadata_path,
        output_dir=args.output_dir,
        study_name=args.study_name,
        formats=args.formats,
    )
    for p in paths:
        logger.info("Created: %s", p)
